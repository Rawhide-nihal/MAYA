"""MAYA universal attachment intelligence.

Parses supported local files into structured, bounded context. The analyzer reports
capabilities/confidence explicitly and never pretends image semantics were extracted
when no visual model/OCR is available.
"""
from __future__ import annotations

import csv
import json
import mimetypes
import os
import re
import zipfile
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, List, Optional

from PIL import Image
from agents.vision.ocr import windows_ocr


TEXT_EXTENSIONS = {
    ".txt", ".md", ".rst", ".log", ".ini", ".cfg", ".conf", ".toml", ".yaml", ".yml",
    ".json", ".jsonl", ".xml", ".html", ".css", ".js", ".jsx", ".ts", ".tsx",
    ".py", ".java", ".c", ".h", ".cpp", ".hpp", ".cs", ".go", ".rs", ".php",
    ".sh", ".bash", ".ps1", ".bat", ".cmd", ".sql", ".gradle", ".properties",
}

PROJECT_MARKERS = {
    "package.json", "requirements.txt", "pyproject.toml", "setup.py", "pom.xml",
    "build.gradle", "build.gradle.kts", "cargo.toml", "go.mod", "composer.json",
    "dockerfile", "docker-compose.yml", "docker-compose.yaml", ".gitignore",
}

LANGUAGE_BY_EXT = {
    ".py": "Python", ".js": "JavaScript", ".jsx": "JavaScript/React",
    ".ts": "TypeScript", ".tsx": "TypeScript/React", ".java": "Java",
    ".c": "C", ".h": "C/C++ headers", ".cpp": "C++", ".hpp": "C++",
    ".cs": "C#", ".go": "Go", ".rs": "Rust", ".php": "PHP",
    ".html": "HTML", ".css": "CSS", ".sql": "SQL", ".sh": "Shell",
    ".ps1": "PowerShell",
}


class AttachmentIntelligence:
    def __init__(self, max_text_chars: int = 24000, max_zip_entries: int = 5000):
        self.max_text_chars = max_text_chars
        self.max_zip_entries = max_zip_entries

    def analyze(self, filepath: str) -> Dict[str, Any]:
        path = Path(filepath).expanduser().resolve()
        if not path.exists() or not path.is_file():
            return {"success": False, "error": f"Attachment not found: {path}"}

        ext = path.suffix.lower()
        base = {
            "success": True,
            "path": str(path),
            "name": path.name,
            "type": ext.lstrip(".") or "unknown",
            "mime_type": mimetypes.guess_type(str(path))[0],
            "size_bytes": path.stat().st_size,
            "confidence": "high",
            "warnings": [],
        }

        try:
            if ext == ".pdf":
                details = self._pdf(path)
            elif ext == ".docx":
                details = self._docx(path)
            elif ext == ".pptx":
                details = self._pptx(path)
            elif ext == ".xlsx":
                details = self._xlsx(path)
            elif ext == ".csv":
                details = self._csv(path)
            elif ext == ".zip":
                details = self._zip(path)
            elif ext in TEXT_EXTENSIONS:
                details = self._text(path)
            elif ext in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif", ".tiff"}:
                details = self._image(path)
            else:
                return {
                    **base,
                    "success": False,
                    "confidence": "none",
                    "error": f"Unsupported attachment type: {ext or 'no extension'}",
                    "capabilities": [],
                }
        except Exception as exc:
            return {
                **base,
                "success": False,
                "confidence": "low",
                "error": f"Failed to analyze {path.name}: {exc}",
            }

        result = {**base, **details}
        result["context_text"] = self._build_context_text(result)
        return result

    def _bounded(self, text: str) -> str:
        if len(text) <= self.max_text_chars:
            return text
        return text[: self.max_text_chars] + "\n...[content truncated by MAYA context budget]"

    def _analyze_image_blob(self, blob: bytes, label: str) -> Dict[str, Any]:
        try:
            with Image.open(BytesIO(blob)) as image:
                width, height = image.size
                fmt = image.format
                ocr = windows_ocr.recognize_image(image.copy())
            text = str(ocr.get("text") or "").strip() if ocr.get("success") else ""
            return {
                "label": label,
                "width": width,
                "height": height,
                "format": fmt,
                "ocr_text": text[:5000],
                "ocr_success": bool(ocr.get("success")),
            }
        except Exception as exc:
            return {
                "label": label,
                "ocr_success": False,
                "error": str(exc),
            }

    def _text(self, path: Path) -> Dict[str, Any]:
        raw = path.read_text(encoding="utf-8", errors="replace")
        lines = raw.splitlines()
        return {
            "capabilities": ["text_extraction", "line_analysis", "code_analysis"],
            "line_count": len(lines),
            "extracted_text": self._bounded(raw),
            "summary": f"Text/code file with {len(lines)} lines.",
        }

    def _pdf(self, path: Path) -> Dict[str, Any]:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        pages: List[Dict[str, Any]] = []
        combined = []
        embedded_images = []

        for index, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            pages.append({
                "page": index,
                "text_preview": text[:1200],
                "char_count": len(text),
            })
            if text:
                combined.append(f"\n--- PAGE {index} ---\n{text}")

            if len(embedded_images) < 20:
                try:
                    for image_index, image_file in enumerate(page.images):
                        if len(embedded_images) >= 20:
                            break
                        embedded_images.append(
                            self._analyze_image_blob(
                                image_file.data,
                                f"page {index} image {image_index + 1}: {getattr(image_file, 'name', '')}"
                            )
                        )
                except Exception:
                    pass

        tables = []
        try:
            import pdfplumber
            with pdfplumber.open(str(path)) as pdf:
                for page_index, page in enumerate(pdf.pages[:80], start=1):
                    try:
                        found = page.extract_tables() or []
                    except Exception:
                        found = []
                    for table_index, table in enumerate(found[:20], start=1):
                        rows = []
                        for row in (table or [])[:200]:
                            rows.append([
                                "" if cell is None else str(cell)
                                for cell in (row or [])[:40]
                            ])
                        if rows:
                            tables.append({
                                "page": page_index,
                                "table": table_index,
                                "rows": rows,
                            })
        except Exception:
            # PDF text extraction still works when optional table heuristics fail.
            pass

        for table in tables[:40]:
            combined.append(
                f"\n--- PAGE {table['page']} TABLE {table['table']} ---"
            )
            combined.extend(" | ".join(row) for row in table["rows"][:100])

        for image in embedded_images:
            if image.get("ocr_text"):
                combined.append(
                    f"\n--- EMBEDDED IMAGE OCR {image.get('label')} ---\n"
                    + image["ocr_text"]
                )

        extracted = self._bounded("".join(combined))
        warnings = []
        if not extracted.strip():
            warnings.append(
                "No extractable PDF text/table/OCR content was found. The PDF may contain visual material "
                "that is not readable by the current grounded local extractors."
            )
        if embedded_images and not any(img.get("ocr_text") for img in embedded_images):
            warnings.append(
                "Embedded PDF images were detected, but no readable OCR text was recovered from them. "
                "MAYA will not invent their visual meaning."
            )

        return {
            "capabilities": [
                "page_text_extraction", "page_cross_reference",
                "table_extraction", "embedded_image_metadata", "embedded_image_ocr"
            ],
            "page_count": len(reader.pages),
            "pages": pages[:80],
            "table_count": len(tables),
            "tables": tables[:40],
            "embedded_image_count_analyzed": len(embedded_images),
            "embedded_images": embedded_images,
            "extracted_text": extracted,
            "warnings": warnings,
            "confidence": "high" if extracted.strip() else "low",
            "summary": (
                f"PDF with {len(reader.pages)} page(s), {len(tables)} extracted table(s), "
                f"and {len(embedded_images)} embedded image(s) analyzed; "
                f"text recovered from {sum(1 for p in pages if p['char_count'])} page(s)."
            ),
        }

    def _docx(self, path: Path) -> Dict[str, Any]:
        from docx import Document

        doc = Document(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        tables = []
        for t_index, table in enumerate(doc.tables, start=1):
            rows = []
            for row in table.rows[:200]:
                rows.append([cell.text for cell in row.cells[:40]])
            tables.append({"table": t_index, "rows": rows})

        embedded_images = []
        try:
            with zipfile.ZipFile(path, "r") as archive:
                media = [
                    name for name in archive.namelist()
                    if name.lower().startswith("word/media/")
                ][:30]
                for name in media:
                    embedded_images.append(
                        self._analyze_image_blob(archive.read(name), name)
                    )
        except Exception:
            pass

        text_parts = paragraphs[:]
        for table in tables:
            text_parts.append(f"\n--- TABLE {table['table']} ---")
            text_parts.extend(" | ".join(row) for row in table["rows"])
        for image in embedded_images:
            if image.get("ocr_text"):
                text_parts.append(
                    f"\n--- EMBEDDED IMAGE OCR {image.get('label')} ---\n"
                    + image["ocr_text"]
                )

        warnings = []
        if embedded_images and not any(img.get("ocr_text") for img in embedded_images):
            warnings.append(
                "Embedded Word images were detected, but no readable OCR text was recovered; "
                "their visual meaning is not inferred."
            )

        return {
            "capabilities": [
                "paragraph_extraction", "table_extraction",
                "embedded_image_metadata", "embedded_image_ocr"
            ],
            "paragraph_count": len(paragraphs),
            "table_count": len(tables),
            "tables": tables[:30],
            "embedded_image_count": len(embedded_images),
            "embedded_images": embedded_images,
            "warnings": warnings,
            "extracted_text": self._bounded("\n".join(text_parts)),
            "summary": (
                f"Word document with {len(paragraphs)} non-empty paragraph(s), "
                f"{len(tables)} table(s), and {len(embedded_images)} embedded image(s)."
            ),
        }

    def _pptx(self, path: Path) -> Dict[str, Any]:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        prs = Presentation(str(path))
        slides = []
        combined = []
        embedded_images = []
        for idx, slide in enumerate(prs.slides, start=1):
            texts = []
            slide_images = []
            for shape_index, shape in enumerate(slide.shapes, start=1):
                text = getattr(shape, "text", "")
                if text and text.strip():
                    texts.append(text.strip())

                if (
                    getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE
                    and len(embedded_images) < 40
                ):
                    try:
                        image = self._analyze_image_blob(
                            shape.image.blob,
                            f"slide {idx} picture {shape_index}"
                        )
                        embedded_images.append(image)
                        slide_images.append(image)
                    except Exception:
                        pass

            slides.append({
                "slide": idx,
                "text": texts,
                "images": slide_images,
            })
            combined.append(f"\n--- SLIDE {idx} ---\n" + "\n".join(texts))
            for image in slide_images:
                if image.get("ocr_text"):
                    combined.append(
                        f"\n[Slide {idx} image OCR] {image['ocr_text']}"
                    )

        warnings = []
        if embedded_images and not any(img.get("ocr_text") for img in embedded_images):
            warnings.append(
                "PowerPoint pictures were detected, but no readable OCR text was recovered; "
                "MAYA will not invent their visual meaning."
            )

        return {
            "capabilities": [
                "slide_text_extraction", "slide_cross_reference",
                "embedded_image_metadata", "embedded_image_ocr"
            ],
            "slide_count": len(prs.slides),
            "slides": slides,
            "embedded_image_count": len(embedded_images),
            "embedded_images": embedded_images,
            "warnings": warnings,
            "extracted_text": self._bounded("".join(combined)),
            "summary": (
                f"PowerPoint with {len(prs.slides)} slide(s) and "
                f"{len(embedded_images)} picture(s) analyzed."
            ),
        }

    def _xlsx(self, path: Path) -> Dict[str, Any]:
        import openpyxl

        workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=False)
        sheets = []
        combined = []
        for sheet in workbook.worksheets:
            rows = []
            for r_idx, row in enumerate(sheet.iter_rows(values_only=False), start=1):
                if r_idx > 250:
                    break
                values = []
                for cell in row[:30]:
                    if cell.value is None:
                        values.append("")
                    elif isinstance(cell.value, str) and cell.value.startswith("="):
                        values.append(cell.value)
                    else:
                        values.append(str(cell.value))
                if any(v for v in values):
                    rows.append(values)
            sheets.append({"name": sheet.title, "rows": rows})
            combined.append(f"\n--- SHEET {sheet.title} ---")
            combined.extend(" | ".join(row) for row in rows[:120])

        workbook.close()
        return {
            "capabilities": ["sheet_extraction", "table_extraction", "formula_visibility"],
            "sheet_count": len(sheets),
            "sheets": sheets,
            "extracted_text": self._bounded("\n".join(combined)),
            "summary": f"Excel workbook with {len(sheets)} sheet(s).",
        }

    def _csv(self, path: Path) -> Dict[str, Any]:
        rows = []
        with path.open("r", encoding="utf-8", errors="replace", newline="") as handle:
            reader = csv.reader(handle)
            for index, row in enumerate(reader):
                if index >= 500:
                    break
                rows.append(row)

        width = max((len(row) for row in rows), default=0)
        text = "\n".join(" | ".join(row) for row in rows)
        return {
            "capabilities": ["table_extraction", "row_column_analysis"],
            "row_count_sampled": len(rows),
            "column_count_max": width,
            "rows": rows[:200],
            "extracted_text": self._bounded(text),
            "summary": f"CSV with {len(rows)} sampled row(s), up to {width} column(s).",
        }

    def _image(self, path: Path) -> Dict[str, Any]:
        with Image.open(path) as image:
            width, height = image.size
            mode = image.mode
            fmt = image.format

        ocr = windows_ocr.recognize_file(str(path))
        extracted = str(ocr.get("text") or "").strip() if ocr.get("success") else ""
        capabilities = ["image_metadata", "visual_context_reference"]
        warnings = []
        confidence = "medium"

        if extracted:
            capabilities.extend(["windows_ocr", "text_bounding_boxes"])
            confidence = "high"
        else:
            warnings.append(
                "No readable OCR text was extracted. MAYA will retain/reference the image but will not invent visual details "
                "that are not available through OCR or another grounded vision source."
            )

        return {
            "capabilities": capabilities,
            "width": width,
            "height": height,
            "image_mode": mode,
            "image_format": fmt,
            "ocr": ocr,
            "extracted_text": self._bounded(extracted),
            "confidence": confidence,
            "warnings": warnings,
            "summary": (
                f"Image {width}x{height} ({fmt or path.suffix}); "
                f"OCR extracted {len(extracted)} character(s)."
            ),
        }

    def _zip(self, path: Path) -> Dict[str, Any]:
        with zipfile.ZipFile(path, "r") as archive:
            entries = [i for i in archive.infolist() if not i.is_dir()]
            original_entry_count = len(entries)
            if len(entries) > self.max_zip_entries:
                entries = entries[: self.max_zip_entries]
                truncated = True
            else:
                truncated = False

            names = [i.filename.replace("\\", "/") for i in entries]
            languages: Dict[str, int] = {}
            markers = []
            source_entries = []
            config_entries = []
            source_relationships: Dict[str, List[str]] = {}

            for name in names:
                p = Path(name)
                ext = p.suffix.lower()
                lang = LANGUAGE_BY_EXT.get(ext)
                if lang:
                    languages[lang] = languages.get(lang, 0) + 1
                    source_entries.append(name)
                if p.name.lower() in PROJECT_MARKERS:
                    markers.append(name)
                if p.name.lower() in {
                    "package.json", "requirements.txt", "pyproject.toml", "pom.xml",
                    "cargo.toml", "go.mod", "composer.json", "build.gradle", "build.gradle.kts"
                }:
                    config_entries.append(name)

            dependencies: Dict[str, List[str]] = {}
            config_summaries: Dict[str, Any] = {}
            excerpts = []
            total_excerpt_chars = 0

            for info in entries:
                name = info.filename.replace("\\", "/")
                p = Path(name)
                ext = p.suffix.lower()
                lower_name = p.name.lower()

                if info.file_size > 1_000_000:
                    continue

                is_text = ext in TEXT_EXTENSIONS or lower_name in PROJECT_MARKERS
                if not is_text:
                    continue

                try:
                    text = archive.read(info).decode("utf-8", errors="replace")
                except Exception:
                    continue

                # Parse dependency/config files into structured project metadata.
                try:
                    if lower_name == "package.json":
                        data = json.loads(text)
                        deps = []
                        for key in ("dependencies", "devDependencies", "peerDependencies", "optionalDependencies"):
                            section = data.get(key, {})
                            if isinstance(section, dict):
                                deps.extend(f"{pkg}@{version}" for pkg, version in section.items())
                        dependencies[name] = sorted(set(deps))
                        config_summaries[name] = {
                            "name": data.get("name"),
                            "version": data.get("version"),
                            "scripts": data.get("scripts", {}),
                        }
                    elif lower_name == "requirements.txt":
                        deps = [
                            line.strip() for line in text.splitlines()
                            if line.strip() and not line.lstrip().startswith("#")
                        ]
                        dependencies[name] = deps[:300]
                    elif lower_name in {"pyproject.toml", "cargo.toml"}:
                        try:
                            import tomllib
                            data = tomllib.loads(text)
                            if lower_name == "pyproject.toml":
                                project = data.get("project", {}) if isinstance(data, dict) else {}
                                deps = list(project.get("dependencies", []) or [])
                                optional = project.get("optional-dependencies", {}) or {}
                                if isinstance(optional, dict):
                                    for group, values in optional.items():
                                        deps.extend(f"[{group}] {value}" for value in (values or []))
                                dependencies[name] = [str(x) for x in deps][:300]
                                config_summaries[name] = {
                                    "project_name": project.get("name"),
                                    "requires_python": project.get("requires-python"),
                                }
                            else:
                                deps_section = data.get("dependencies", {}) if isinstance(data, dict) else {}
                                if isinstance(deps_section, dict):
                                    dependencies[name] = [
                                        f"{pkg}={value}" for pkg, value in deps_section.items()
                                    ][:300]
                        except Exception:
                            pass
                    elif lower_name == "go.mod":
                        deps = []
                        in_block = False
                        for raw in text.splitlines():
                            line = raw.strip()
                            if line.startswith("require ("):
                                in_block = True
                                continue
                            if in_block and line == ")":
                                in_block = False
                                continue
                            if line.startswith("require "):
                                deps.append(line[len("require "):].strip())
                            elif in_block and line and not line.startswith("//"):
                                deps.append(line)
                        dependencies[name] = deps[:300]
                    elif lower_name == "pom.xml":
                        import xml.etree.ElementTree as ET
                        root = ET.fromstring(text)
                        deps = []
                        for dep in root.findall(".//{*}dependency"):
                            group = dep.findtext("{*}groupId") or ""
                            artifact = dep.findtext("{*}artifactId") or ""
                            version = dep.findtext("{*}version") or ""
                            if artifact:
                                value = f"{group}:{artifact}" if group else artifact
                                if version:
                                    value += f":{version}"
                                deps.append(value)
                        dependencies[name] = deps[:300]
                    elif lower_name == "composer.json":
                        data = json.loads(text)
                        deps = []
                        for key in ("require", "require-dev"):
                            section = data.get(key, {})
                            if isinstance(section, dict):
                                deps.extend(f"{pkg}@{version}" for pkg, version in section.items())
                        dependencies[name] = sorted(set(deps))
                    elif lower_name in {"build.gradle", "build.gradle.kts"}:
                        deps = re.findall(
                            r"(?:implementation|api|compileOnly|runtimeOnly|testImplementation)\s*\(?[\"']([^\"']+)",
                            text
                        )
                        dependencies[name] = deps[:300]
                except Exception:
                    # Config parsing must never make archive analysis fail.
                    pass

                # Lightweight source relationship extraction: imports/requires only.
                if ext == ".py":
                    refs = re.findall(
                        r"^\s*(?:from\s+([a-zA-Z0-9_\.]+)\s+import|import\s+([a-zA-Z0-9_\.]+))",
                        text,
                        flags=re.MULTILINE
                    )
                    flat = [a or b for a, b in refs if a or b]
                    if flat:
                        source_relationships[name] = sorted(set(flat))[:80]
                elif ext in {".js", ".jsx", ".ts", ".tsx"}:
                    refs = re.findall(
                        r"(?:from\s+[\"']([^\"']+)[\"']|require\(\s*[\"']([^\"']+)[\"']\s*\))",
                        text
                    )
                    flat = [a or b for a, b in refs if a or b]
                    if flat:
                        source_relationships[name] = sorted(set(flat))[:80]
                elif ext == ".java":
                    refs = re.findall(r"^\s*import\s+([a-zA-Z0-9_\.]+);", text, flags=re.MULTILINE)
                    if refs:
                        source_relationships[name] = sorted(set(refs))[:80]

                if total_excerpt_chars < self.max_text_chars:
                    chunk = f"\n--- FILE {name} ---\n{text[:5000]}"
                    excerpts.append(chunk)
                    total_excerpt_chars += len(chunk)

        top_dirs = sorted({name.split("/", 1)[0] for name in names if "/" in name})[:60]
        language_list = sorted(languages.items(), key=lambda kv: kv[1], reverse=True)
        warnings = []
        if truncated:
            warnings.append(
                f"ZIP contained {original_entry_count} files; analysis was bounded to the first {self.max_zip_entries} entries."
            )

        dependency_count = sum(len(v) for v in dependencies.values())
        return {
            "capabilities": [
                "archive_structure", "project_structure", "dependency_config_detection",
                "dependency_extraction", "source_language_detection",
                "source_relationship_detection", "bounded_source_extraction"
            ],
            "entry_count": original_entry_count,
            "analyzed_entry_count": len(names),
            "entries": names[:1000],
            "top_level_directories": top_dirs,
            "project_markers": markers,
            "config_files": config_entries,
            "config_summaries": config_summaries,
            "dependencies": dependencies,
            "dependency_count": dependency_count,
            "source_relationships": source_relationships,
            "languages": [{"language": lang, "files": count} for lang, count in language_list],
            "source_file_count": len(source_entries),
            "extracted_text": self._bounded("".join(excerpts)),
            "warnings": warnings,
            "summary": (
                f"ZIP/project archive with {original_entry_count} file(s), "
                f"{len(source_entries)} recognized source file(s), "
                f"{len(markers)} project/dependency marker(s), and "
                f"{dependency_count} dependency declaration(s)."
            ),
        }

    def _build_context_text(self, result: Dict[str, Any]) -> str:
        parts = [
            f"Attachment: {result.get('name')}",
            f"Type: {result.get('type')}",
            f"Summary: {result.get('summary', '')}",
            f"Confidence: {result.get('confidence', 'unknown')}",
        ]

        warnings = result.get("warnings") or []
        if warnings:
            parts.append("Warnings: " + " | ".join(str(w) for w in warnings))

        if result.get("project_markers"):
            parts.append("Project markers: " + ", ".join(result["project_markers"][:20]))
        if result.get("languages"):
            parts.append(
                "Languages: " + ", ".join(
                    f"{item['language']} ({item['files']})" for item in result["languages"][:12]
                )
            )

        dependencies = result.get("dependencies") or {}
        if dependencies:
            dep_lines = []
            for config_name, values in list(dependencies.items())[:12]:
                dep_lines.append(f"{config_name}: " + ", ".join(str(v) for v in values[:40]))
            parts.append("Dependencies:\n" + "\n".join(dep_lines))

        relationships = result.get("source_relationships") or {}
        if relationships:
            rel_lines = []
            for source_name, refs in list(relationships.items())[:40]:
                rel_lines.append(f"{source_name} -> " + ", ".join(str(v) for v in refs[:30]))
            parts.append("Source relationships:\n" + "\n".join(rel_lines))

        extracted = result.get("extracted_text") or ""
        if extracted:
            parts.append("Extracted content:\n" + self._bounded(extracted))

        return "\n".join(parts)
